#!/usr/bin/env python3
"""Docs check: automatic consistency check of the repository documents.

This is not a sign-off. Only the owner verifies and promotes work to uat and main.

Run from anywhere:  python scripts/verify_docs.py
Add --stamp to record the result line in docs/handover/project-memory.md.
Exit code is 0 only if every check passes.
"""
import glob
import os
import re
import subprocess
import sys
import tempfile
import shutil
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
results = []


def check(name, cond, detail=""):
    results.append(bool(cond))
    print(("PASS " if cond else "FAIL ") + name + ((" -> " + str(detail)) if detail and not cond else ""))


def read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


def norm(p):
    return p.replace("\\", "/")


mds = [norm(f) for f in glob.glob("**/*.md", recursive=True)] + [norm(f) for f in glob.glob(".github/*.md")]
mds = sorted(set(f for f in mds if not f.startswith("node_modules/")))

# 1. relative links resolve
bad = []
for f in mds:
    for m in re.finditer(r"\]\(([^)#\s]+)(#[^)]*)?\)", read(f)):
        link = m.group(1)
        if link.startswith(("http://", "https://", "mailto:")):
            continue
        if not os.path.exists(os.path.normpath(os.path.join(os.path.dirname(f), link))):
            bad.append((f, link))
check("relative links resolve", not bad, bad)

# 2. stale names or paths (superseded reference docs are exempt)
stale = []
patterns = [r"`backlog\.md`", r"\bKaki\b", r"backlog/index\.html", r"backlog/data\.js", r"kaki_data_plan",
            r"vision_estimate", r"clinically_validated", r"causal_evidence"]
for f in mds:
    if f.startswith("docs/reference/"):
        continue
    t = read(f)
    for p in patterns:
        if f.endswith("handover/project-memory.md") and "Kaki" in p:
            continue  # the decisions register records the rename
        if re.search(p, t) and not (f.endswith("specs/memory.md") and p in (r"clinically_validated", r"causal_evidence")):
            stale.append((f, p))
check("no stale names or paths", not stale, stale)

# 3. secrets and connection strings
secret_patterns = [r"sb_publishable_", r"sb_secret_", r"eyJhbGci", r"sk-[A-Za-z0-9]{20,}", r"AKIA[0-9A-Z]{12,}",
                   r"service_role\s*[:=]\s*\S{20,}", r"postgres(ql)?://[^\s:@/]+:[^\s@]+@"]
hits = []
for f in glob.glob("**/*", recursive=True) + glob.glob(".github/*") + glob.glob(".githooks/*"):
    f = norm(f)
    if os.path.isdir(f) or f.endswith((".xlsx", ".png", ".jpg", ".pdf")) or f.startswith(("data/raw/", "data/processed/", "node_modules/")):
        continue
    try:
        t = read(f)
    except Exception:
        continue
    if f == "scripts/verify_docs.py":
        continue
    for p in secret_patterns:
        if re.search(p, t):
            hits.append((f, p))
check("no keys or connection strings in tracked files", not hits, hits)

# 4. backlog integrity (three personal backlogs plus the generated tracker)
people_files = ["docs/backlog/hongyi.md", "docs/backlog/joely.md", "docs/backlog/damien.md"]
heading = re.compile(r"^### (\S+) (.+?) · (P\d)( · NEEDS OK)?\s*$", re.M)
all_ids, blocks, owner_of = [], [], {}
for pf in people_files:
    bt = read(pf)
    ms = list(heading.finditer(bt))
    for i, m in enumerate(ms):
        body = bt[m.end():(ms[i + 1].start() if i + 1 < len(ms) else len(bt))]
        all_ids.append(m.group(1))
        owner_of[m.group(1)] = pf
        blocks.append((m.group(1), body))
check("backlog task ids unique across the three backlogs", len(all_ids) == len(set(all_ids)), len(all_ids))
check("every task has Status, Needs first, Needed later, Note and an Ask of 20+ characters",
      all(re.search(r"^- \*\*Status:\*\* (backlog|in progress|in review|blocked|done)$", x, re.M)
          and "- **Needs first:**" in x and "- **Needed later:**" in x and "- **Note:**" in x
          and re.search(r"^- \*\*Ask:\*\* .{20,}", x, re.M) for _, x in blocks))
bad_dep = []
for tid, blk in blocks:
    for key in ("Needs first", "Needed later"):
        m = re.search(r"^- \*\*%s:\*\* (.*)$" % key, blk, re.M)
        if m and m.group(1).strip().lower() != "none":
            bad_dep += [(tid, d.strip()) for d in m.group(1).split(",") if d.strip() not in all_ids]
check("every dependency exists", not bad_dep, bad_dep)
check("every task has subtasks", all(re.search(r"^  - \[.\] ", x, re.M) for _, x in blocks))
check("each person has a BLOCKERS block", all("<!-- BLOCKERS:START -->" in read(pf) for pf in people_files))
done_wrong = []
for tid, blk in blocks:
    status = re.search(r"^- \*\*Status:\*\* (.+)$", blk, re.M).group(1)
    subs = re.findall(r"^  - \[(.)\] ", blk, re.M)
    if status == "done" and any(s != "x" for s in subs):
        done_wrong.append((tid, "done with unticked subtasks"))
    if subs and all(s == "x" for s in subs) and status in ("backlog", "in progress", "blocked"):
        done_wrong.append((tid, "every subtask ticked but Status is " + status))
check("Status agrees with the subtask boxes", not done_wrong, done_wrong)
tr = subprocess.run([sys.executable, "scripts/update_tracker.py", "--check"], capture_output=True, text=True)
check("generated tracker (SOT.md and BLOCKERS blocks) is up to date", tr.returncode == 0, (tr.stdout + tr.stderr).strip())

# 5. schema layer map covers every table
sch = read("docs/specs/schema.md")
lm = sch[sch.index("## 0b. Layer map"):sch.index("## 1. Enums")]
layer_tables = set(re.findall(r"`([a-z_]+)`", lm))
heads = set(re.findall(r"^### ([a-z_]+)", sch, re.M)) | set(re.findall(r"`([a-z_]+)\(id", sch))
check("every schema table is in the layer map", all(h in layer_tables for h in heads), [h for h in heads if h not in layer_tables])

# 6. session logs and project memory
mem_path = "docs/handover/project-memory.md"
mem = read(mem_path) if os.path.exists(mem_path) else ""
check("project memory exists", bool(mem))
logs = [os.path.basename(f) for f in glob.glob("docs/handover/sessions/*.md") if os.path.basename(f).lower() != "readme.md"]
bad_names = [l for l in logs if not re.match(r"^[a-z][a-z0-9-]*_\d{4}-\d{2}-\d{2}_\d{4}\.md$", l)]
check("session log names are <user>_<YYYY-MM-DD_HHMM>.md", not bad_names, bad_names)
missing_index = [l for l in logs if l[:-3] not in mem]
check("every session log is listed in project-memory.md", not missing_index, missing_index)

# 7. workbook (optional dependency)
xlsx = "data/plan/smart_health_data_plan.xlsx"
try:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    check("data plan workbook has its 3 sheets", wb.sheetnames == ["Use now", "Still missing", "Domain knowledge"], wb.sheetnames)
except ImportError:
    print("SKIP data plan workbook check (openpyxl not installed)")
except FileNotFoundError:
    check("data plan workbook exists", False, xlsx)

# 8. layout
required = ["README.md", ".gitignore", ".gitattributes", ".github/pull_request_template.md", ".github/CODEOWNERS",
            ".githooks/pre-commit", ".githooks/pre-push", "docs/specs/prd.md", "docs/specs/architecture.md",
            "docs/specs/memory.md", "docs/specs/schema.md", "docs/specs/flow.md", "docs/specs/design.md",
            "docs/specs/rules.md", "docs/backlog/SOT.md", "docs/backlog/hongyi.md", "docs/backlog/joely.md",
            "docs/backlog/damien.md", "docs/backlog/details/e1-data-foundation.md", "docs/handover/HANDOVER.md",
            "docs/handover/session-log-template.md", "docs/handover/project-memory.md", "docs/mockup/mockup-prompt.md",
            "data/README.md", "scripts/verify_docs.py", "scripts/update_tracker.py"]
check("required files present", all(os.path.exists(p) for p in required), [p for p in required if not os.path.exists(p)])
root_md = [f for f in os.listdir(".") if f.endswith(".md")]
check("only README.md at the repo root", root_md == ["README.md"], root_md)

# 9. dev-only hooks behave (temporary repo, this repo's config is not touched)
if shutil.which("git") and shutil.which("sh"):
    tmp = tempfile.mkdtemp()
    hooks = norm(os.path.abspath(".githooks"))

    def git(*a, env=None):
        e = dict(os.environ)
        e.update(env or {})
        return subprocess.run(["git"] + list(a), cwd=tmp, capture_output=True, text=True, env=e)

    git("init", "-q", "-b", "main")
    git("config", "user.email", "t@t")
    git("config", "user.name", "t")
    check("hook blocks commits on main", git("-c", "core.hooksPath=" + hooks, "commit", "--allow-empty", "-m", "x").returncode != 0)
    check("owner override allows a commit on main", git("-c", "core.hooksPath=" + hooks, "commit", "--allow-empty", "-m", "x", env={"ALLOW_PROMOTE": "1"}).returncode == 0)
    git("checkout", "-q", "-b", "dev")
    check("hook allows commits on dev", git("-c", "core.hooksPath=" + hooks, "commit", "--allow-empty", "-m", "y").returncode == 0)

    def push(ref, env=None):
        e = dict(os.environ)
        e.update(env or {})
        return subprocess.run(["sh", ".githooks/pre-push", "origin", "url"], input="refs/heads/dev abc %s def\n" % ref,
                              capture_output=True, text=True, env=e).returncode

    check("pre-push blocks main and uat", push("refs/heads/main") != 0 and push("refs/heads/uat") != 0)
    check("pre-push allows dev and feature branches", push("refs/heads/dev") == 0 and push("refs/heads/feature/e9-2-x") == 0)
    shutil.rmtree(tmp, ignore_errors=True)
else:
    print("SKIP hook checks (git or sh not available)")

passed = sum(results)
total = len(results)
ok = passed == total
stamp = "%s: %s (%d of %d checks passed)" % (datetime.now().strftime("%Y-%m-%d %H:%M"), "PASS" if ok else "FAIL", passed, total)
print("\nOVERALL: " + stamp)

if "--stamp" in sys.argv and os.path.exists(mem_path):
    t = read(mem_path)
    t = re.sub(r"^\*\*Last docs check:\*\*.*$", "**Last docs check:** " + stamp, t, count=1, flags=re.M)
    with open(mem_path, "w", encoding="utf-8", newline="\n") as f:
        f.write(t)
    print("Stamped " + mem_path)

sys.exit(0 if ok else 1)
