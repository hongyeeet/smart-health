#!/usr/bin/env python3
"""Measure the medicine scanner against a ground-truth sheet you wrote (prd F2, docs/specs/f2-scan.md).

Usage:
  python scripts/scan_accuracy.py data/raw/label_photos /path/to/medication_ground_truth.xlsx
  python scripts/scan_accuracy.py data/raw/label_photos truth.xlsx --name "Your Name"   also black out lines with your own name
  python scripts/scan_accuracy.py ... --only 1 3                                         run only those photo numbers

The photos (sorted by file name) are matched to the sheet rows in order: photo 1 is row 1, and so on. The sheet columns are
photo, item_type, name, strength, no_of_tabs, times_per_day, when_to_take, meal_relation, drowsiness, total_no_of_tabs, start_date.
Dates in the sheet are read as day-month-year; a date a spreadsheet turned into a real date has its day and month swapped back.

For every field it reports: correct or not, the confidence the scanner gave, and whether the 0.85 rule would have made the user
type it. The summary counts how often a RIGHT value was sent to typing (extra work for the user) and how often a WRONG value was
trusted (the dangerous case). Nothing is saved to the database. The report goes to <photos>/accuracy_report.csv (git-ignored folder).
"""
import argparse
import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "services" / "api"))
sys.path.insert(0, str(ROOT / "scripts"))

from app.agents.scan_extractor import extract_scan  # noqa: E402
from app.core.settings import get_settings  # noqa: E402
from app.llm.vertex import VertexClient  # noqa: E402
from app.services import scan_accuracy as acc  # noqa: E402
from app.services.scan_privacy import default_ocr  # noqa: E402
from app.services.scan_store import load_drug_index  # noqa: E402


def read_sheet(path: Path) -> list[dict]:
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    from openpyxl import load_workbook
    ws = load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h).strip() for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(c is not None for c in r)]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("photos")
    ap.add_argument("truth")
    ap.add_argument("--name", action="append", default=[])
    ap.add_argument("--only", nargs="*", type=int)
    a = ap.parse_args()
    s = get_settings()
    photos = sorted(p for p in Path(a.photos).iterdir() if p.suffix.lower() in (".jpg", ".jpeg", ".png"))
    rows = read_sheet(Path(a.truth))
    if len(photos) != len(rows):
        print(f"warning: {len(photos)} photos but {len(rows)} sheet rows; matching the first {min(len(photos), len(rows))} in order")

    import psycopg
    from check_supabase import clean_url
    with psycopg.connect(clean_url(s.database_url.get_secret_value()), connect_timeout=20) as conn:
        index = load_drug_index(conn)
    client = VertexClient(s.google_api_key.get_secret_value(), s.google_cloud_project, s.google_cloud_region)
    ocr = default_ocr()

    report = []
    for n, (photo, row) in enumerate(zip(photos, rows), 1):
        if a.only and n not in a.only:
            continue
        exp = acc.expected_from_sheet(row)
        outcome, _ = extract_scan(photo.read_bytes(), client=client, model=s.llm_vision_model, index=index, ocr=ocr,
                                  known_names=a.name, threshold=s.scan_confidence_threshold)
        print(f"\nphoto {n}: status={outcome.status} review_mode={outcome.review_mode} reasons={outcome.reasons} masked={outcome.mask_summary}")
        item = outcome.items[0] if outcome.items else None
        for field in acc.COMPARED:
            if field == "item_type":
                got, conf, flagged = (item.item_type if item else None), (1.0 if item else 0.0), False
            else:
                f = item.fields.get(field) if item else None
                got = acc.normalise_scanned(field, f.value if f else None)
                conf, flagged = (f.confidence if f else 0.0), bool(f and f.needs_confirmation)
            ok = acc.same(field, exp[field], got)
            report.append(dict(photo=n, field=field, expected=exp[field], got=got, confidence=round(conf, 2), flagged=flagged, correct=ok,
                               has_value=got is not None))
            mark = "ok   " if ok else "WRONG"
            print(f"  {mark} {field:<20} expected={exp[field]!s:<30} got={got!s:<30} conf={conf:.2f}{'  TYPE IT' if flagged else ''}")

    out = Path(a.photos) / "accuracy_report.csv"
    with out.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(report[0]))
        w.writeheader()
        w.writerows(report)

    print("\n=== per field (all photos) ===")
    print(f"{'field':<20}{'correct':>9}{'right but sent to typing':>26}{'wrong but trusted':>20}")
    tot = [0, 0, 0, 0]
    for field in acc.COMPARED:
        r = [x for x in report if x["field"] == field]
        c = sum(x["correct"] for x in r)
        flagged_right = sum(x["correct"] and x["flagged"] and x["has_value"] for x in r)
        trusted_wrong = sum((not x["correct"]) and x["has_value"] and not x["flagged"] for x in r)
        print(f"{field:<20}{c:>5}/{len(r):<3}{flagged_right:>26}{trusted_wrong:>20}")
        tot[0] += c; tot[1] += len(r); tot[2] += flagged_right; tot[3] += trusted_wrong
    print(f"{'ALL':<20}{tot[0]:>5}/{tot[1]:<3}{tot[2]:>26}{tot[3]:>20}")
    print(f"\nreport written to {out}")


if __name__ == "__main__":
    main()
